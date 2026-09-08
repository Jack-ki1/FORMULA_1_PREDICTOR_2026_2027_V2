"""
CSV and Excel report generator.
Exports predictions in CSV and Excel formats.
"""
import csv
import io
from typing import Dict, List, Any
import json
from datetime import datetime


def extract_predictions_list(data: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Extract a flat list of predictions from either list or target dictionary."""
    raw = data.get('predictions', [])
    if isinstance(raw, list):
        return raw
    if isinstance(raw, dict):
        target_id = data.get('target_id')
        if target_id and target_id in raw:
            target_obj = raw[target_id]
            if isinstance(target_obj, dict):
                return target_obj.get('predictions', [])
            elif isinstance(target_obj, list):
                return target_obj
        # Fallback to first target
        for _, v in raw.items():
            if isinstance(v, dict) and 'predictions' in v:
                return v['predictions']
            elif isinstance(v, list):
                return v
    return []


class CSVExcelReportGenerator:
    """Generates CSV and Excel reports from prediction data."""
    
    def generate_csv(self, data: Dict[str, Any]) -> str:
        """
        Generate CSV report from prediction data.
        
        Args:
            data: Prediction data dictionary
        
        Returns:
            CSV string
        """
        output = io.StringIO()
        
        predictions = extract_predictions_list(data)
        if not predictions:
            return "No predictions available"
        
        writer = csv.writer(output)
        writer.writerow(['Driver', 'Probability', 'Percentage', 'Rank'])
        
        for i, prediction in enumerate(predictions, 1):
            driver_code = prediction.get('driver_code', 'Unknown')
            probability = prediction.get('probability', 0.0)
            percentage = prediction.get('percentage', probability * 100)
            
            writer.writerow([driver_code, f"{probability:.4f}", f"{percentage:.1f}%", i])
        
        writer.writerow([])
        writer.writerow(['Target', data.get('target_id', 'Unknown')])
        writer.writerow(['Race ID', data.get('race_id', 'Unknown')])
        writer.writerow(['Session', data.get('session', data.get('session_type', 'Unknown'))])
        writer.writerow(['Sub Session', data.get('sub_session', '')])
        writer.writerow(['Confidence', f"{data.get('confidence', 0.0):.2f}"])
        
        return output.getvalue()
    
    def generate_excel(self, data: Dict[str, Any]) -> bytes:
        """Generate Excel report from prediction data."""
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Predictions"
            
            headers = ['Driver', 'Probability', 'Percentage', 'Rank']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="E10600", end_color="E10600", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            predictions = extract_predictions_list(data)
            for i, prediction in enumerate(predictions, 2):
                driver_code = prediction.get('driver_code', 'Unknown')
                probability = prediction.get('probability', 0.0)
                percentage = prediction.get('percentage', probability * 100)
                
                ws.cell(row=i, column=1, value=driver_code)
                ws.cell(row=i, column=2, value=f"{probability:.4f}")
                ws.cell(row=i, column=3, value=f"{percentage:.1f}%")
                ws.cell(row=i, column=4, value=i - 1)
            
            metadata_row = len(predictions) + 3
            ws.cell(row=metadata_row, column=1, value="Metadata", font=Font(bold=True))
            
            metadata = [
                ['Target', data.get('target_id', 'Unknown')],
                ['Race ID', data.get('race_id', 'Unknown')],
                ['Session', data.get('session', data.get('session_type', 'Unknown'))],
                ['Confidence', f"{data.get('confidence', 0.0):.2f}"],
            ]
            
            for i, (key, value) in enumerate(metadata, metadata_row + 1):
                ws.cell(row=i, column=1, value=key)
                ws.cell(row=i, column=2, value=value)
            
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output.read()
            
        except ImportError:
            csv_data = self.generate_csv(data)
            return csv_data.encode('utf-8')
    
    def generate_json(self, data: Dict[str, Any]) -> Dict[str, Any]:
        """Generate JSON report from prediction data."""
        return {
            'report_type': 'predictions',
            'generated_at': datetime.utcnow().isoformat(),
            'data': data,
        }
    
    def generate_summary_text(self, data: Dict[str, Any]) -> str:
        """Generate text summary from prediction data."""
        predictions = extract_predictions_list(data)
        if not predictions:
            return "No predictions available"
        
        summary = f"F1 Predictor 2026 - Race Predictions\n"
        summary += f"Race: {data.get('race_id', 'Unknown')}\n"
        summary += f"Session: {data.get('session', data.get('session_type', 'Unknown'))}\n"
        summary += f"Target: {data.get('target_id', 'Unknown')}\n\n"
        
        summary += "Top Predictions:\n"
        for i, prediction in enumerate(predictions[:5], 1):
            driver_code = prediction.get('driver_code', 'Unknown')
            probability = prediction.get('probability', 0.0)
            percentage = prediction.get('percentage', probability * 100)
            summary += f"{i}. {driver_code}: {percentage:.1f}%\n"
        
        return summary
